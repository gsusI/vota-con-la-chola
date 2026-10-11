#!/usr/bin/env python3
"""Independently validate the complete real-document native-text corpus."""

from __future__ import annotations

import argparse
import codecs
import gzip
import hashlib
import json
import os
import re
import sqlite3
import sys
import tempfile
import uuid
import xml.etree.ElementTree as ET
from collections import Counter
from collections.abc import Iterator, Mapping
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from datetime import UTC, date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import ZipFile

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from publicdata_docs.extraction import (
    DEFAULT_MAX_XLSX_TEXT_CHARS,
    DOCUMENT_EXTRACTION_ALGORITHM_VERSION,
)
from publicdata_docs.generation_manifest import (
    GenerationManifestError,
    iter_generation_rows,
)

DEFAULT_RAW_ROOT = Path("etl/data/raw")
DEFAULT_TEXT_ROOT = Path("etl/data/derived/real-document-text-corpus-v2")
DEFAULT_INVENTORY = Path("etl/data/manifests/real-document-format-inventory.jsonl")
DEFAULT_PROVENANCE = Path("etl/data/manifests/real-document-provenance-files.jsonl")
DEFAULT_MANIFEST = Path("etl/data/published/real-document-text-manifest-latest.jsonl")
DEFAULT_EXTRACTION_REPORT = Path(
    "docs/etl/sprints/SCALE-FOUNDATION-20260810/evidence/"
    "real-document-text-extraction.json"
)
DEFAULT_VALIDATION_REPORT = Path(
    "docs/etl/sprints/SCALE-FOUNDATION-20260810/evidence/"
    "real-document-text-validation.json"
)
SHA256_HEX_LENGTH = 64
_XLSX_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
XLSX_VALIDATION_MAX_IN_FLIGHT_PER_WORKER = 2


class CorpusValidationError(RuntimeError):
    """Raised when extracted corpus evidence fails exact validation."""


def utc_now() -> str:
    return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def display_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(REPO_ROOT.resolve()).as_posix()
    except ValueError:
        return resolved.name


def safe_child(root: Path, relative: str) -> Path:
    if not relative or Path(relative).is_absolute():
        raise CorpusValidationError(f"unsafe relative path: {relative!r}")
    candidate = (root / relative).resolve()
    try:
        candidate.relative_to(root.resolve())
    except ValueError as exc:
        raise CorpusValidationError(f"path escapes root: {relative}") from exc
    return candidate


def sha256_path(path: Path, chunk_bytes: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_bytes), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-root", default=str(DEFAULT_RAW_ROOT))
    parser.add_argument("--text-root", default=str(DEFAULT_TEXT_ROOT))
    parser.add_argument("--inventory", default=str(DEFAULT_INVENTORY))
    parser.add_argument("--provenance", default=str(DEFAULT_PROVENANCE))
    parser.add_argument("--manifest", default=str(DEFAULT_MANIFEST))
    parser.add_argument("--extraction-report", default=str(DEFAULT_EXTRACTION_REPORT))
    parser.add_argument("--report-out", default=str(DEFAULT_VALIDATION_REPORT))
    parser.add_argument("--workers", type=int, default=12)
    parser.add_argument("--batch-size", type=int, default=128)
    parser.add_argument(
        "--full-raw-check",
        action="store_true",
        help="Hash every file instance, including duplicate-content paths.",
    )
    parser.add_argument("--enforce", action="store_true")
    return parser.parse_args(argv)


def iter_jsonl(path: Path) -> Iterator[dict[str, Any]]:
    if not path.is_file():
        raise CorpusValidationError(f"missing JSONL input: {display_path(path)}")
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            try:
                item = json.loads(line)
            except json.JSONDecodeError as exc:
                raise CorpusValidationError(
                    f"invalid JSONL at {display_path(path)}:{line_number}: {exc}"
                ) from exc
            if not isinstance(item, dict):
                raise CorpusValidationError(
                    f"JSONL row must be an object: {display_path(path)}:{line_number}"
                )
            yield item


def load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise CorpusValidationError(f"missing JSON report: {display_path(path)}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise CorpusValidationError(
            f"invalid JSON report {display_path(path)}: {exc}"
        ) from exc
    if not isinstance(payload, dict):
        raise CorpusValidationError("extraction report must be a JSON object")
    return payload


def open_index() -> tuple[sqlite3.Connection, Path]:
    file_descriptor, name = tempfile.mkstemp(
        prefix="vota-real-document-text-validation-", suffix=".db"
    )
    os.close(file_descriptor)
    path = Path(name)
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA synchronous = OFF")
    conn.execute("PRAGMA temp_store = FILE")
    conn.executescript(
        """
        CREATE TABLE expected (
          path TEXT PRIMARY KEY,
          content_sha256 TEXT NOT NULL,
          bytes INTEGER NOT NULL,
          provenance_seen INTEGER NOT NULL DEFAULT 0,
          provenance_status TEXT,
          public_url_edges INTEGER NOT NULL DEFAULT 0,
          ocr_candidate_pages INTEGER NOT NULL DEFAULT 0
        );
        CREATE INDEX idx_expected_sha256 ON expected(content_sha256);

        CREATE TABLE observed (
          path TEXT PRIMARY KEY,
          content_sha256 TEXT NOT NULL,
          bytes INTEGER NOT NULL,
          extension TEXT NOT NULL,
          mime_type TEXT NOT NULL,
          provenance_status TEXT,
          public_url_edges INTEGER NOT NULL,
          ocr_candidate_pages INTEGER NOT NULL,
          work_state TEXT NOT NULL,
          extraction_status TEXT,
          extraction_method TEXT NOT NULL,
          text_sha256 TEXT,
          text_cas_path TEXT,
          text_chars INTEGER,
          text_truncated INTEGER,
          last_error TEXT
        );
        CREATE INDEX idx_observed_sha256 ON observed(content_sha256);
        CREATE INDEX idx_observed_text_sha256 ON observed(text_sha256);
        """
    )
    return conn, path


def load_expected(
    conn: sqlite3.Connection, inventory_path: Path, provenance_path: Path
) -> int:
    rows = 0
    for item in iter_jsonl(inventory_path):
        path = str(item.get("path") or "")
        sha256 = str(item.get("sha256") or "").lower()
        declared_bytes = item.get("bytes")
        if (
            item.get("accountability_eligible") is not True
            or item.get("status") != "ok"
            or len(sha256) != SHA256_HEX_LENGTH
            or not isinstance(declared_bytes, int)
            or declared_bytes < 0
        ):
            raise CorpusValidationError(f"invalid inventory row: {path}")
        page_quality = item.get("page_text_quality") or {}
        candidate_pages = (
            page_quality.get("ocr_candidate_pages")
            if isinstance(page_quality, dict)
            else []
        )
        conn.execute(
            """
            INSERT INTO expected (
              path, content_sha256, bytes, ocr_candidate_pages
            ) VALUES (?, ?, ?, ?)
            """,
            (path, sha256, declared_bytes, len(candidate_pages or [])),
        )
        rows += 1
        if rows % 5_000 == 0:
            conn.commit()
    conn.commit()

    provenance_rows = 0
    for item in iter_jsonl(provenance_path):
        path = str(item.get("path") or "")
        if (
            item.get("checksum_conflict") is not False
            or item.get("declared_bytes_conflict") is not False
        ):
            raise CorpusValidationError(f"conflicted provenance row: {path}")
        cursor = conn.execute(
            """
            UPDATE expected
            SET provenance_seen = 1, provenance_status = ?, public_url_edges = ?
            WHERE path = ? AND content_sha256 = ? AND bytes = ?
              AND provenance_seen = 0
            """,
            (
                str(item.get("provenance_status") or "unlinked"),
                int(item.get("public_url_edges") or 0),
                path,
                str(item.get("sha256") or "").lower(),
                int(item.get("bytes") or 0),
            ),
        )
        if cursor.rowcount != 1:
            raise CorpusValidationError(f"provenance mismatch: {path}")
        provenance_rows += 1
        if provenance_rows % 5_000 == 0:
            conn.commit()
    conn.commit()
    missing = conn.execute(
        "SELECT COUNT(*) FROM expected WHERE provenance_seen = 0"
    ).fetchone()[0]
    if provenance_rows != rows or int(missing):
        raise CorpusValidationError(
            "expected inventory/provenance sets do not reconcile"
        )
    return rows


def load_observed(conn: sqlite3.Connection, manifest_path: Path) -> int:
    rows = 0
    try:
        manifest_rows = iter_generation_rows(
            manifest_path,
            dataset="native_text",
            row_schema_version="real_document_text_manifest_row_v1",
        )
        for item in manifest_rows:
            path = str(item.get("path") or "")
            if (
                item.get("schema_version") != "real_document_text_manifest_row_v1"
                or item.get("synthetic_or_mock") is not False
                or item.get("official_public_domain_personal_information_retained")
                is not True
                or item.get("native_extraction_algorithm_version")
                != DOCUMENT_EXTRACTION_ALGORITHM_VERSION
            ):
                raise CorpusValidationError(f"invalid extraction manifest row: {path}")
            ocr_pages = item.get("ocr_candidate_pages")
            if not isinstance(ocr_pages, list):
                raise CorpusValidationError(f"invalid OCR candidate list: {path}")
            conn.execute(
                """
                INSERT INTO observed (
                  path, content_sha256, bytes, extension, mime_type,
                  provenance_status,
                  public_url_edges, ocr_candidate_pages, work_state,
                  extraction_status, extraction_method, text_sha256,
                  text_cas_path, text_chars, text_truncated, last_error
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    path,
                    str(item.get("content_sha256") or "").lower(),
                    int(item.get("bytes") or 0),
                    str(item.get("extension") or "").casefold().lstrip("."),
                    str(item.get("mime_type") or ""),
                    str(item.get("provenance_status") or ""),
                    int(item.get("public_url_edges") or 0),
                    len(ocr_pages),
                    str(item.get("work_state") or ""),
                    str(item.get("native_extraction_status") or ""),
                    str(item.get("native_extraction_method") or ""),
                    item.get("native_text_sha256"),
                    item.get("native_text_cas_path"),
                    item.get("native_text_chars"),
                    (
                        1
                        if item.get("native_text_truncated") is True
                        else 0
                        if item.get("native_text_truncated") is False
                        else None
                    ),
                    item.get("last_error"),
                ),
            )
            rows += 1
            if rows % 5_000 == 0:
                conn.commit()
    except GenerationManifestError as exc:
        raise CorpusValidationError(f"invalid extraction manifest: {exc}") from exc
    conn.commit()
    return rows


def _mismatch_count(conn: sqlite3.Connection) -> int:
    return int(
        conn.execute(
            """
            SELECT COUNT(*)
            FROM expected e
            LEFT JOIN observed o ON o.path = e.path
            WHERE o.path IS NULL
               OR o.content_sha256 != e.content_sha256
               OR o.bytes != e.bytes
               OR o.provenance_status != e.provenance_status
               OR o.public_url_edges != e.public_url_edges
               OR o.ocr_candidate_pages != e.ocr_candidate_pages
            """
        ).fetchone()[0]
    ) + int(
        conn.execute(
            """
            SELECT COUNT(*)
            FROM observed o
            LEFT JOIN expected e ON e.path = o.path
            WHERE e.path IS NULL
            """
        ).fetchone()[0]
    )


def _verify_raw_item(raw_root: Path, item: Mapping[str, Any]) -> str | None:
    relative = str(item["path"])
    try:
        path = safe_child(raw_root, relative)
        if not path.is_file():
            return f"missing:{relative}"
        if path.stat().st_size != int(item["bytes"]):
            return f"bytes:{relative}"
        if sha256_path(path) != str(item["content_sha256"]):
            return f"sha256:{relative}"
        return None
    except Exception as exc:  # noqa: BLE001
        return f"{type(exc).__name__}:{relative}:{exc}"


def validate_raw_files(
    conn: sqlite3.Connection,
    *,
    raw_root: Path,
    workers: int,
    batch_size: int,
    full_raw_check: bool,
) -> tuple[int, list[str]]:
    if full_raw_check:
        cursor = conn.execute(
            "SELECT path, content_sha256, bytes FROM expected ORDER BY path"
        )
    else:
        cursor = conn.execute(
            """
            SELECT MIN(path) AS path, content_sha256, MIN(bytes) AS bytes
            FROM expected
            GROUP BY content_sha256
            ORDER BY content_sha256
            """
        )
    checked = 0
    errors: list[str] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        while rows := cursor.fetchmany(max(1, batch_size)):
            results = executor.map(
                lambda row: _verify_raw_item(raw_root, dict(row)), rows
            )
            for result in results:
                checked += 1
                if result and len(errors) < 20:
                    errors.append(result)
    return checked, errors


def _verify_text_item(text_root: Path, item: Mapping[str, Any]) -> str | None:
    text_sha256 = str(item["text_sha256"])
    relative = str(item["text_cas_path"])
    expected_relative = (
        Path(text_sha256[:2]) / text_sha256[2:4] / f"{text_sha256}.txt.gz"
    ).as_posix()
    if relative != expected_relative:
        return f"cas_path:{relative}"
    try:
        path = safe_child(text_root, relative)
        if not path.is_file():
            return f"missing:{relative}"
        digest = hashlib.sha256()
        decoder = codecs.getincrementaldecoder("utf-8")("strict")
        chars = 0
        with gzip.open(path, "rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
                chars += len(decoder.decode(chunk))
            chars += len(decoder.decode(b"", final=True))
        if digest.hexdigest() != text_sha256:
            return f"sha256:{relative}"
        if chars != int(item["text_chars"]):
            return f"chars:{relative}:{chars}!={item['text_chars']}"
        return None
    except Exception as exc:  # noqa: BLE001
        return f"{type(exc).__name__}:{relative}:{exc}"


def validate_text_cas(
    conn: sqlite3.Connection,
    *,
    text_root: Path,
    manifest_path: Path,
    workers: int,
    batch_size: int,
) -> tuple[int, int, int, int, int, list[str]]:
    cursor = conn.execute(
        """
        SELECT text_sha256, MIN(text_cas_path) AS text_cas_path,
               MIN(text_chars) AS text_chars,
               MAX(text_chars) AS max_text_chars,
               COUNT(DISTINCT text_cas_path) AS paths
        FROM observed
        WHERE extraction_status = 'native_text_extracted'
        GROUP BY text_sha256
        ORDER BY text_sha256
        """
    )
    checked = 0
    expected_paths: set[str] = set()
    retained_items: dict[str, dict[str, Any]] = {}
    errors: list[str] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        while rows := cursor.fetchmany(max(1, batch_size)):
            items = [dict(row) for row in rows]
            for item in items:
                relative = str(item["text_cas_path"])
                expected_paths.add(relative)
                retained_items[relative] = item
                if (
                    int(item["paths"]) != 1
                    or int(item["text_chars"]) != int(item["max_text_chars"])
                ) and len(errors) < 20:
                    errors.append(f"manifest_conflict:{item['text_sha256']}")
            results = executor.map(
                lambda item: _verify_text_item(text_root, item), items
            )
            for result in results:
                checked += 1
                if result and len(errors) < 20:
                    errors.append(result)
    actual_paths = {
        path.relative_to(text_root).as_posix()
        for path in text_root.rglob("*.txt.gz")
        if path.is_file()
    }
    history_root = (
        manifest_path.parent / f"{manifest_path.stem}-history"
    )
    retained_paths = set(expected_paths)
    retained_generations = 0
    if history_root.is_dir():
        for history_manifest in sorted(history_root.glob("*.json")):
            for item in iter_generation_rows(
                history_manifest,
                dataset="native_text",
                row_schema_version="real_document_text_manifest_row_v1",
                allow_legacy_jsonl=False,
                artifact_root=manifest_path.parent,
            ):
                relative = item.get("native_text_cas_path")
                if relative:
                    relative = str(relative)
                    retained_paths.add(relative)
                    retained_item = {
                        "text_sha256": str(item.get("native_text_sha256") or ""),
                        "text_cas_path": relative,
                        "text_chars": int(item.get("native_text_chars") or 0),
                    }
                    prior = retained_items.get(relative)
                    if prior and (
                        str(prior["text_sha256"]),
                        int(prior["text_chars"]),
                    ) != (
                        retained_item["text_sha256"],
                        retained_item["text_chars"],
                    ):
                        errors.append(f"retained_manifest_conflict:{relative}")
                    retained_items[relative] = retained_item
            retained_generations += 1
    retained_only = [
        retained_items[path] for path in sorted(retained_paths - expected_paths)
    ]
    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        for result in executor.map(
            lambda item: _verify_text_item(text_root, item), retained_only
        ):
            if result and len(errors) < 20:
                errors.append(result)
    extras = sorted(actual_paths - retained_paths)
    missing = sorted(expected_paths - actual_paths)
    retained_missing = sorted(retained_paths - actual_paths)
    for prefix, paths in (
        ("orphan", extras),
        ("missing", missing),
        ("retained_missing", retained_missing),
    ):
        for path in paths[: max(0, 20 - len(errors))]:
            errors.append(f"{prefix}:{path}")
    return (
        checked,
        len(actual_paths),
        len(retained_paths),
        retained_generations,
        len(retained_only),
        errors,
    )


def _independent_rph_items(container: ET.Element) -> list[str]:
    items: list[str] = []
    for child in container:
        if child.tag != _XLSX_MAIN_NS + "rPh":
            continue
        payload = {
            "text": "".join(
                node.text or ""
                for node in child.iter(_XLSX_MAIN_NS + "t")
            ),
            "attributes": sorted(
                (str(key), str(value)) for key, value in child.attrib.items()
            ),
        }
        items.append(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        )
    return items


def _iter_independent_xml_groups(
    handle: Any, *, tag: str
) -> Iterator[ET.Element]:
    """Yield one high-level group at a time without production parser reuse."""

    root: ET.Element | None = None
    for event, element in ET.iterparse(handle, events=("start", "end")):
        if root is None and event == "start":
            root = element
        if event != "end" or element.tag != tag:
            continue
        yield element
        element.clear()
        if root is not None and root is not element:
            root.clear()


def _independent_xlsx_rph_counter(path: Path) -> Counter[str]:
    """Resolve raw shared/inline rPh occurrences without production helpers."""

    shared: dict[int, tuple[str, ...]] = {}
    observed: Counter[str] = Counter()
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        shared_path = "xl/sharedStrings.xml"
        if shared_path in names:
            with archive.open(shared_path) as handle:
                for index, item in enumerate(
                    _iter_independent_xml_groups(
                        handle, tag=_XLSX_MAIN_NS + "si"
                    )
                ):
                    annotations = tuple(_independent_rph_items(item))
                    if annotations:
                        shared[index] = annotations

        for name in sorted(
            value
            for value in names
            if value.startswith("xl/worksheets/") and value.endswith(".xml")
        ):
            with archive.open(name) as handle:
                for row in _iter_independent_xml_groups(
                    handle, tag=_XLSX_MAIN_NS + "row"
                ):
                    for cell in row:
                        if cell.tag != _XLSX_MAIN_NS + "c":
                            continue
                        cell_type = str(cell.get("t") or "n")
                        if cell_type == "s":
                            value_node = cell.find(_XLSX_MAIN_NS + "v")
                            if value_node is None or not value_node.text:
                                continue
                            try:
                                annotations = shared.get(int(value_node.text), ())
                            except ValueError as exc:
                                raise CorpusValidationError(
                                    f"invalid XLSX shared-string index: {path.name}"
                                ) from exc
                            observed.update(annotations)
                        elif cell_type == "inlineStr":
                            inline = cell.find(_XLSX_MAIN_NS + "is")
                            if inline is not None:
                                observed.update(_independent_rph_items(inline))
    return observed


def _cas_phonetic_items(
    raw_cell: Any,
    *,
    relative_path: str,
    row_number: int,
) -> list[str]:
    if not isinstance(raw_cell, list) or len(raw_cell) not in {4, 5}:
        raise CorpusValidationError(
            f"invalid XLSX CAS cell contract: {relative_path}:row={row_number}"
        )
    if len(raw_cell) == 4:
        return []
    encoded = raw_cell[4]
    if not isinstance(encoded, str) or not encoded:
        raise CorpusValidationError(
            f"invalid XLSX CAS phonetic contract: {relative_path}:row={row_number}"
        )
    try:
        annotations = json.loads(encoded)
    except json.JSONDecodeError as exc:
        raise CorpusValidationError(
            f"invalid XLSX CAS phonetic JSON: {relative_path}:row={row_number}"
        ) from exc
    if not isinstance(annotations, list) or not annotations:
        raise CorpusValidationError(
            f"invalid XLSX CAS phonetic annotations: {relative_path}:row={row_number}"
        )
    canonical: list[str] = []
    for item in annotations:
        if (
            not isinstance(item, dict)
            or not isinstance(item.get("text"), str)
            or not isinstance(item.get("attributes"), list)
        ):
            raise CorpusValidationError(
                f"invalid XLSX CAS phonetic item: {relative_path}:row={row_number}"
            )
        attributes = item["attributes"]
        if any(
            not isinstance(pair, list)
            or len(pair) != 2
            or not all(isinstance(value, str) for value in pair)
            for pair in attributes
        ):
            raise CorpusValidationError(
                f"invalid XLSX CAS phonetic attributes: {relative_path}:row={row_number}"
            )
        normalized = {
            "text": item["text"],
            "attributes": sorted((pair[0], pair[1]) for pair in attributes),
        }
        canonical.append(
            json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))
        )
    return canonical


def _cas_cell_semantic(
    raw_cell: Any,
    *,
    relative_path: str,
    row_number: int,
) -> tuple[str, tuple[str, Any], bool]:
    if (
        not isinstance(raw_cell, list)
        or len(raw_cell) not in {4, 5}
        or not all(isinstance(value, str) for value in raw_cell[:4])
    ):
        raise CorpusValidationError(
            f"invalid XLSX CAS cell contract: {relative_path}:row={row_number}"
        )
    coordinate, cell_type, raw_value, raw_formula = raw_cell[:4]
    if raw_formula:
        try:
            formula = json.loads(raw_formula)
        except json.JSONDecodeError as exc:
            raise CorpusValidationError(
                f"invalid XLSX CAS formula JSON: {relative_path}:{coordinate}"
            ) from exc
        if not isinstance(formula, dict) or not isinstance(formula.get("text"), str):
            raise CorpusValidationError(
                f"invalid XLSX CAS formula contract: {relative_path}:{coordinate}"
            )
        return coordinate, ("formula", formula["text"]), True
    if cell_type == "b":
        if raw_value not in {"0", "1", "false", "true"}:
            raise CorpusValidationError(
                f"invalid XLSX boolean scalar: {relative_path}:{coordinate}"
            )
        return coordinate, ("boolean", raw_value in {"1", "true"}), False
    if cell_type == "n":
        try:
            value = _excel_semantic_number(raw_value)
        except InvalidOperation:
            return coordinate, ("number_lexeme", raw_value), False
        return coordinate, ("number", value), False
    if cell_type == "d":
        return coordinate, ("date", raw_value), False
    if cell_type == "e":
        return coordinate, ("error", raw_value), False
    return coordinate, ("text", raw_value), False


def _excel_semantic_number(value: Any) -> Decimal:
    number = Decimal(str(value))
    if not number.is_finite():
        return number
    try:
        return Decimal(format(float(number), ".15g"))
    except (OverflowError, ValueError):
        return number.normalize()


def _openpyxl_cell_semantic(cell: Any) -> tuple[tuple[str, Any], bool] | None:
    value = cell.value
    if value is None:
        return None
    if cell.data_type == "f":
        formula = str(value)
        return ("formula", formula[1:] if formula.startswith("=") else formula), True
    if value == "":
        return None
    if isinstance(value, bool) or cell.data_type == "b":
        return ("boolean", bool(value)), False
    if isinstance(value, (datetime, date, time)) or cell.data_type == "d":
        normalized = value.isoformat() if hasattr(value, "isoformat") else str(value)
        return ("date", normalized), False
    if cell.data_type == "e":
        return ("error", str(value)), False
    if isinstance(value, (int, float, Decimal)):
        try:
            return ("number", _excel_semantic_number(value)), False
        except InvalidOperation:
            return ("number_lexeme", str(value)), False
    return ("text", str(value)), False


def _iter_openpyxl_semantic_rows(
    workbook: Any,
) -> Iterator[tuple[int, str, int, list[tuple[str, tuple[str, Any]]], int]]:
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        first_row = int(worksheet.min_row or 1)
        for row_number, cells in enumerate(
            worksheet.iter_rows(min_row=first_row), start=first_row
        ):
            semantic_cells: list[tuple[str, tuple[str, Any]]] = []
            formula_cells = 0
            for cell in cells:
                semantic = _openpyxl_cell_semantic(cell)
                if semantic is None:
                    continue
                value, is_formula = semantic
                semantic_cells.append((str(cell.coordinate), value))
                formula_cells += int(is_formula)
            if semantic_cells:
                yield (
                    sheet_index,
                    str(worksheet.title),
                    row_number,
                    semantic_cells,
                    formula_cells,
                )


def _verify_xlsx_raw_to_cas(
    item: Mapping[str, Any],
    *,
    raw_root: Path,
    text_root: Path,
) -> dict[str, Any]:
    relative_path = str(item["representative_path"])
    content_sha256 = str(item["content_sha256"])
    result: dict[str, Any] = {
        "content_sha256": content_sha256,
        "rows": 0,
        "cells": 0,
        "formula_cells": 0,
        "phonetic_annotations_raw": 0,
        "phonetic_annotations_cas": 0,
        "worksheets": 0,
        "error": None,
    }
    workbook = None
    try:
        if str(item["extraction_method"]) != "xlsx_openxml_path":
            raise CorpusValidationError(
                f"unsupported XLSX extraction method: {relative_path}"
            )
        raw_path = safe_child(raw_root, relative_path)
        if (
            not raw_path.is_file()
            or raw_path.stat().st_size != int(item["bytes"])
            or sha256_path(raw_path) != content_sha256
        ):
            raise CorpusValidationError(
                f"XLSX representative checksum drifted: {relative_path}"
            )
        extraction_status = str(item["extraction_status"])
        if extraction_status not in {"native_text_extracted", "native_no_text"}:
            raise CorpusValidationError(
                f"unsupported XLSX extraction status: {relative_path}"
            )
        raw_phonetic = _independent_xlsx_rph_counter(raw_path)
        result["phonetic_annotations_raw"] = sum(raw_phonetic.values())

        from openpyxl import load_workbook

        workbook = load_workbook(
            raw_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        result["worksheets"] = len(workbook.worksheets)
        raw_rows = _iter_openpyxl_semantic_rows(workbook)
        if extraction_status == "native_no_text":
            if (
                item.get("text_sha256") is not None
                or item.get("text_cas_path") is not None
                or int(item["text_chars"]) != 0
                or next(raw_rows, None) is not None
                or raw_phonetic
            ):
                raise CorpusValidationError(
                    f"XLSX native-no-text contract drifted: {relative_path}"
                )
            return result

        if (
            len(str(item.get("text_sha256") or "")) != SHA256_HEX_LENGTH
            or not str(item.get("text_cas_path") or "")
            or int(item["text_chars"]) < 1
        ):
            raise CorpusValidationError(
                f"XLSX native-text metadata drifted: {relative_path}"
            )
        cas_path = safe_child(text_root, str(item["text_cas_path"]))
        if not cas_path.is_file():
            raise CorpusValidationError(f"XLSX CAS missing: {relative_path}")

        cas_phonetic: Counter[str] = Counter()
        cas_digest = hashlib.sha256()
        cas_chars = 0
        with gzip.open(cas_path, "rt", encoding="utf-8", newline="") as handle:
            for cas_row_index, physical_line in enumerate(handle, start=1):
                has_newline = physical_line.endswith("\n")
                encoded_line = physical_line.encode("utf-8")
                cas_digest.update(encoded_line)
                cas_chars += len(physical_line)
                line = physical_line[:-1] if has_newline else physical_line
                if not line:
                    raise CorpusValidationError(
                        f"empty XLSX CAS row: {relative_path}:line={cas_row_index}"
                    )
                try:
                    cas_row = json.loads(line)
                except json.JSONDecodeError as exc:
                    raise CorpusValidationError(
                        f"invalid XLSX CAS JSON: {relative_path}:line={cas_row_index}"
                    ) from exc
                if (
                    not isinstance(cas_row, dict)
                    or not isinstance(cas_row.get("sheet"), str)
                    or not isinstance(cas_row.get("row"), int)
                    or not isinstance(cas_row.get("cells"), list)
                ):
                    raise CorpusValidationError(
                        "invalid XLSX CAS row contract: "
                        f"{relative_path}:line={cas_row_index}"
                    )
                sheet_name = str(cas_row["sheet"])
                source_row_number = int(cas_row["row"])
                expected_cells: list[tuple[str, tuple[str, Any]]] = []
                expected_formula_cells = 0
                previous_column_index = 0
                for raw_cell in cas_row["cells"]:
                    phonetic_items = _cas_phonetic_items(
                        raw_cell,
                        relative_path=relative_path,
                        row_number=source_row_number,
                    )
                    cas_phonetic.update(phonetic_items)
                    coordinate, semantic, is_formula = _cas_cell_semantic(
                        raw_cell,
                        relative_path=relative_path,
                        row_number=source_row_number,
                    )
                    coordinate_match = re.fullmatch(
                        r"(?P<column>[A-Z]+)(?P<row>[1-9][0-9]*)", coordinate
                    )
                    if (
                        coordinate_match is None
                        or int(coordinate_match.group("row")) != source_row_number
                    ):
                        raise CorpusValidationError(
                            "XLSX CAS coordinate-row mismatch: "
                            f"{relative_path}:{sheet_name}!{coordinate}"
                        )
                    column_index = 0
                    for character in coordinate_match.group("column"):
                        column_index = column_index * 26 + (
                            ord(character) - ord("A") + 1
                        )
                    if (
                        column_index <= previous_column_index
                        or column_index > 16_384
                    ):
                        raise CorpusValidationError(
                            "XLSX CAS column-order/bound mismatch: "
                            f"{relative_path}:{sheet_name}!{coordinate}"
                        )
                    previous_column_index = column_index
                    if not raw_cell[2] and not raw_cell[3] and not phonetic_items:
                        raise CorpusValidationError(
                            "XLSX CAS contains an empty cell: "
                            f"{relative_path}:{sheet_name}!{coordinate}"
                        )
                    if raw_cell[2] or raw_cell[3]:
                        expected_cells.append((coordinate, semantic))
                    expected_formula_cells += int(is_formula)
                if expected_cells:
                    try:
                        (
                            _sheet_index,
                            actual_sheet_name,
                            actual_row_number,
                            actual_cells,
                            actual_formula_cells,
                        ) = next(raw_rows)
                    except StopIteration as exc:
                        raise CorpusValidationError(
                            "XLSX CAS contains an extra semantic row: "
                            f"{relative_path}:line={cas_row_index}"
                        ) from exc
                    if (
                        sheet_name != actual_sheet_name
                        or source_row_number != actual_row_number
                    ):
                        raise CorpusValidationError(
                            "XLSX sheet/row mismatch: "
                            f"{relative_path}:line={cas_row_index}"
                        )
                    if len(expected_cells) != len(actual_cells):
                        raise CorpusValidationError(
                            "XLSX row cell-count mismatch: "
                            f"{relative_path}:{sheet_name}!{source_row_number}"
                        )
                    for expected_cell, actual_cell in zip(
                        expected_cells, actual_cells, strict=True
                    ):
                        if expected_cell[0] != actual_cell[0]:
                            raise CorpusValidationError(
                                "XLSX coordinate mismatch: "
                                f"{relative_path}:{sheet_name}!{source_row_number}"
                            )
                        if expected_cell[1] != actual_cell[1]:
                            raise CorpusValidationError(
                                "XLSX scalar mismatch: "
                                f"{relative_path}:{sheet_name}!{expected_cell[0]}"
                            )
                    if expected_formula_cells != actual_formula_cells:
                        raise CorpusValidationError(
                            "XLSX formula-count mismatch: "
                            f"{relative_path}:{sheet_name}!{source_row_number}"
                        )
                result["rows"] += 1
                result["cells"] += len(cas_row["cells"])
                result["formula_cells"] += expected_formula_cells
        try:
            extra_row = next(raw_rows)
        except StopIteration:
            extra_row = None
        if extra_row is not None:
            raise CorpusValidationError(
                f"XLSX raw workbook contains an extra row: {relative_path}"
            )
        if cas_phonetic != raw_phonetic:
            raise CorpusValidationError(
                f"XLSX phonetic annotation mismatch: {relative_path}"
            )
        result["phonetic_annotations_cas"] = sum(cas_phonetic.values())
        if cas_digest.hexdigest() != str(item["text_sha256"]):
            raise CorpusValidationError(f"XLSX CAS checksum drifted: {relative_path}")
        if cas_chars != int(item["text_chars"]):
            raise CorpusValidationError(
                f"XLSX CAS character count drifted: {relative_path}"
            )
    except Exception as exc:  # noqa: BLE001
        result["error"] = f"{type(exc).__name__}: {exc}"[:2_000]
    finally:
        if workbook is not None:
            workbook.close()
    return result


def validate_xlsx_raw_to_cas(
    conn: sqlite3.Connection,
    *,
    raw_root: Path,
    text_root: Path,
    workers: int,
) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise CorpusValidationError(
            "openpyxl is required for independent XLSX raw-to-CAS validation"
        ) from exc

    cursor = conn.execute(
        """
        SELECT content_sha256,
               MIN(path) AS representative_path,
               MIN(bytes) AS bytes,
               MAX(bytes) AS max_bytes,
               MIN(mime_type) AS mime_type,
               MAX(mime_type) AS max_mime_type,
               MIN(extraction_method) AS extraction_method,
               MAX(extraction_method) AS max_extraction_method,
               MIN(extraction_status) AS extraction_status,
               MAX(extraction_status) AS max_extraction_status,
               MIN(text_sha256) AS text_sha256,
               MAX(text_sha256) AS max_text_sha256,
               MIN(text_cas_path) AS text_cas_path,
               MAX(text_cas_path) AS max_text_cas_path,
               MIN(text_chars) AS text_chars,
               MAX(text_chars) AS max_text_chars,
               COUNT(*) AS file_instances
        FROM observed
        WHERE extension = 'xlsx'
        GROUP BY content_sha256
        ORDER BY content_sha256
        """
    )
    xlsx_workers = min(2, max(1, int(workers)))
    fetch_batch_rows = (
        xlsx_workers * XLSX_VALIDATION_MAX_IN_FLIGHT_PER_WORKER
    )
    distinct_contents = 0
    successful_contents = 0
    worksheets_replayed = 0
    rows_replayed = 0
    cells_replayed = 0
    formula_cells = 0
    phonetic_annotations_observed = 0
    phonetic_annotations_replayed = 0
    failures: list[str] = []
    with ThreadPoolExecutor(max_workers=xlsx_workers) as executor:
        while batch := cursor.fetchmany(fetch_batch_rows):
            pending: dict[Future[dict[str, Any]], str] = {}
            for row in batch:
                distinct_contents += 1
                item = dict(row)
                if (
                    int(item["bytes"]) != int(item["max_bytes"])
                    or str(item["mime_type"]) != str(item["max_mime_type"])
                    or str(item["extraction_method"])
                    != str(item["max_extraction_method"])
                    or str(item["extraction_status"])
                    != str(item["max_extraction_status"])
                    or item["text_sha256"] != item["max_text_sha256"]
                    or item["text_cas_path"] != item["max_text_cas_path"]
                    or int(item["text_chars"]) != int(item["max_text_chars"])
                ):
                    if len(failures) < 20:
                        failures.append(
                            "XLSX duplicate-content manifest drift: "
                            f"{item['content_sha256']}"
                        )
                    continue
                future = executor.submit(
                    _verify_xlsx_raw_to_cas,
                    item,
                    raw_root=raw_root,
                    text_root=text_root,
                )
                pending[future] = str(item["content_sha256"])

            while pending:
                completed, _ = wait(pending, return_when=FIRST_COMPLETED)
                for future in completed:
                    content_sha256 = pending.pop(future)
                    outcome = future.result()
                    if outcome["error"]:
                        if len(failures) < 20:
                            failures.append(
                                f"{content_sha256}: {outcome['error']}"
                            )
                        continue
                    successful_contents += 1
                    worksheets_replayed += int(outcome["worksheets"])
                    rows_replayed += int(outcome["rows"])
                    cells_replayed += int(outcome["cells"])
                    formula_cells += int(outcome["formula_cells"])
                    phonetic_annotations_observed += int(
                        outcome["phonetic_annotations_raw"]
                    )
                    phonetic_annotations_replayed += int(
                        outcome["phonetic_annotations_cas"]
                    )
    return {
        "validator": "openpyxl_semantic_plus_independent_openxml_phonetic_replay_v2",
        "validator_version": str(openpyxl.__version__),
        "workers": xlsx_workers,
        "fetch_batch_rows": fetch_batch_rows,
        "max_in_flight": fetch_batch_rows,
        "file_instances": int(
            conn.execute(
                "SELECT COUNT(*) FROM observed WHERE extension = 'xlsx'"
            ).fetchone()[0]
        ),
        "distinct_content_objects": distinct_contents,
        "distinct_content_objects_replayed": successful_contents,
        "worksheets_replayed": worksheets_replayed,
        "nonempty_rows_replayed": rows_replayed,
        "nonempty_cells_replayed": cells_replayed,
        "formula_cells_observed": formula_cells,
        "phonetic_annotations_observed": phonetic_annotations_observed,
        "phonetic_annotations_replayed": phonetic_annotations_replayed,
        "formula_validation_status": (
            "not_observed_in_real_cohort"
            if formula_cells == 0
            else "formula_text_only_cached_values_not_independently_validated"
        ),
        "failures": failures,
    }


def write_json_atomic(path: Path, payload: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        with temp_path.open("w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=True, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_path, path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def validate(args: argparse.Namespace) -> dict[str, Any]:
    raw_root = Path(args.raw_root)
    text_root = Path(args.text_root)
    inventory_path = Path(args.inventory)
    provenance_path = Path(args.provenance)
    manifest_path = Path(args.manifest)
    extraction_report_path = Path(args.extraction_report)
    extraction_report = load_json(extraction_report_path)
    if (
        extraction_report.get("schema_version") != "real_document_text_extraction_v1"
        or extraction_report.get("status") != "ok"
    ):
        raise CorpusValidationError("extraction report is not a complete v1 run")
    policy = extraction_report.get("policy") or {}
    if (
        policy.get("synthetic_or_mock_records_counted") is not False
        or policy.get("official_public_domain_personal_information_retained")
        is not True
        or policy.get("native_text_is_not_sanitized_or_identity_stripped") is not True
        or policy.get("pdf_plain_text_angle_brackets_retained") is not True
        or policy.get("xlsx_public_phonetic_annotations_retained") is not True
        or extraction_report.get("extraction_algorithm_version")
        != DOCUMENT_EXTRACTION_ALGORITHM_VERSION
        or int(
            (extraction_report.get("effective_configuration") or {}).get(
                "max_text_chars", 0
            )
        )
        != DEFAULT_MAX_XLSX_TEXT_CHARS
    ):
        raise CorpusValidationError("extraction report policy is unsafe")

    conn, index_path = open_index()
    try:
        expected_rows = load_expected(conn, inventory_path, provenance_path)
        observed_rows = load_observed(conn, manifest_path)
        mismatch_rows = _mismatch_count(conn)
        invalid_states = int(
            conn.execute(
                """
                SELECT COUNT(*) FROM observed
                WHERE work_state != 'succeeded'
                   OR extension = ''
                   OR mime_type = ''
                   OR extraction_method = ''
                   OR extraction_status NOT IN (
                     'native_text_extracted', 'native_no_text'
                   )
                   OR last_error IS NOT NULL
                   OR (
                     extraction_status = 'native_text_extracted'
                     AND (
                       LENGTH(text_sha256) != 64
                       OR text_cas_path IS NULL
                       OR text_chars IS NULL
                       OR text_chars <= 0
                     )
                   )
                   OR (
                     extraction_status = 'native_no_text'
                     AND (
                       text_sha256 IS NOT NULL
                       OR text_cas_path IS NOT NULL
                       OR COALESCE(text_chars, 0) != 0
                     )
                   )
                """
            ).fetchone()[0]
        )
        raw_checked, raw_errors = validate_raw_files(
            conn,
            raw_root=raw_root,
            workers=int(args.workers),
            batch_size=int(args.batch_size),
            full_raw_check=bool(args.full_raw_check),
        )
        (
            text_checked,
            actual_cas_files,
            retained_cas_files,
            retained_native_generations,
            retained_superseded_cas_checked,
            text_errors,
        ) = validate_text_cas(
            conn,
            text_root=text_root,
            manifest_path=manifest_path,
            workers=int(args.workers),
            batch_size=int(args.batch_size),
        )
        xlsx_replay = validate_xlsx_raw_to_cas(
            conn,
            raw_root=raw_root,
            text_root=text_root,
            workers=int(args.workers),
        )
        totals_row = conn.execute(
            """
            SELECT
              COUNT(DISTINCT content_sha256) AS distinct_content_objects,
              SUM(CASE WHEN public_url_edges > 0 THEN 1 ELSE 0 END)
                AS public_url_instances,
              SUM(CASE WHEN public_url_edges = 0 THEN 1 ELSE 0 END)
                AS provenance_gap_instances,
              SUM(ocr_candidate_pages) AS ocr_candidate_pages
            FROM observed
            """
        ).fetchone()
        extraction_row = conn.execute(
            """
            SELECT
              COUNT(*) AS distinct_content_objects,
              SUM(CASE WHEN extraction_status = 'native_no_text' THEN 1 ELSE 0 END)
                AS no_text_objects,
              SUM(CASE WHEN text_truncated = 1 THEN 1 ELSE 0 END)
                AS truncated_objects,
              SUM(COALESCE(text_chars, 0)) AS native_text_chars
            FROM (
              SELECT content_sha256, MAX(extraction_status) AS extraction_status,
                     MAX(text_truncated) AS text_truncated,
                     MAX(text_chars) AS text_chars
              FROM observed
              GROUP BY content_sha256
            )
            """
        ).fetchone()
    finally:
        conn.close()
        index_path.unlink(missing_ok=True)
        index_path.with_name(f"{index_path.name}-wal").unlink(missing_ok=True)
        index_path.with_name(f"{index_path.name}-shm").unlink(missing_ok=True)

    report_manifest = extraction_report.get("manifest") or {}
    report_totals = extraction_report.get("totals") or {}
    manifest_sha256 = sha256_path(manifest_path)
    totals = {
        "file_instances": observed_rows,
        "distinct_content_objects": int(totals_row["distinct_content_objects"]),
        "public_url_instances": int(totals_row["public_url_instances"] or 0),
        "provenance_gap_instances": int(totals_row["provenance_gap_instances"] or 0),
        "native_text_chars": int(extraction_row["native_text_chars"] or 0),
        "native_no_text_objects": int(extraction_row["no_text_objects"] or 0),
        "native_truncated_objects": int(extraction_row["truncated_objects"] or 0),
        "ocr_candidate_pages": int(totals_row["ocr_candidate_pages"] or 0),
        "raw_files_checksum_verified": raw_checked,
        "text_cas_objects_checksum_verified": text_checked,
        "text_cas_files_observed": text_checked,
        "text_cas_files_observed_across_retained_generations": actual_cas_files,
        "text_cas_files_expected_across_retained_generations": retained_cas_files,
        "retained_native_generation_manifests_checked": retained_native_generations,
        "retained_superseded_text_cas_objects_checksum_verified": (
            retained_superseded_cas_checked
        ),
        "xlsx_file_instances": int(xlsx_replay["file_instances"]),
        "xlsx_distinct_content_objects": int(
            xlsx_replay["distinct_content_objects"]
        ),
        "xlsx_distinct_content_objects_raw_to_cas_replayed": int(
            xlsx_replay["distinct_content_objects_replayed"]
        ),
        "xlsx_worksheets_raw_to_cas_replayed": int(
            xlsx_replay["worksheets_replayed"]
        ),
        "xlsx_nonempty_rows_raw_to_cas_replayed": int(
            xlsx_replay["nonempty_rows_replayed"]
        ),
        "xlsx_nonempty_cells_raw_to_cas_replayed": int(
            xlsx_replay["nonempty_cells_replayed"]
        ),
        "xlsx_formula_cells_observed_in_real_cohort": int(
            xlsx_replay["formula_cells_observed"]
        ),
        "xlsx_phonetic_annotations_observed_in_real_cohort": int(
            xlsx_replay["phonetic_annotations_observed"]
        ),
        "xlsx_phonetic_annotations_raw_to_cas_replayed": int(
            xlsx_replay["phonetic_annotations_replayed"]
        ),
    }
    reconciled_report_keys = (
        "file_instances",
        "distinct_content_objects",
        "public_url_instances",
        "provenance_gap_instances",
        "native_text_chars",
        "native_no_text_objects",
        "native_truncated_objects",
        "ocr_candidate_pages",
    )
    report_totals_match = all(
        int(report_totals.get(key, -1)) == int(totals[key])
        for key in reconciled_report_keys
    )
    checks = {
        "expected_and_observed_rows_match": expected_rows == observed_rows,
        "manifest_rows_match_extraction_report": (
            int(report_manifest.get("rows", -1)) == observed_rows
        ),
        "manifest_sha256_matches_extraction_report": (
            report_manifest.get("sha256") == manifest_sha256
        ),
        "inventory_provenance_and_manifest_values_match": mismatch_rows == 0,
        "all_work_states_are_successful_and_well_formed": invalid_states == 0,
        "raw_checksums_match": not raw_errors,
        "text_cas_checksums_and_character_counts_match": not text_errors,
        "text_cas_inventory_is_exact_across_retained_generations": (
            retained_cas_files == actual_cas_files
        ),
        "extraction_report_totals_reconcile": report_totals_match,
        "no_native_text_truncation": totals["native_truncated_objects"] == 0,
        "no_synthetic_or_mock_records": True,
        "official_public_domain_personal_information_retained": True,
        "extraction_algorithm_version_pinned_to_report_and_manifest": (
            extraction_report.get("extraction_algorithm_version")
            == DOCUMENT_EXTRACTION_ALGORITHM_VERSION
        ),
        "real_xlsx_cohort_present": int(xlsx_replay["file_instances"]) > 0,
        "every_distinct_xlsx_raw_to_cas_replayed": (
            int(xlsx_replay["distinct_content_objects"]) > 0
            and int(xlsx_replay["distinct_content_objects_replayed"])
            == int(xlsx_replay["distinct_content_objects"])
        ),
        "xlsx_raw_to_cas_semantic_replay_has_no_failures": not xlsx_replay[
            "failures"
        ],
        "xlsx_formula_validation_not_overclaimed": (
            int(xlsx_replay["formula_cells_observed"]) == 0
            and xlsx_replay["formula_validation_status"]
            == "not_observed_in_real_cohort"
        ),
        "xlsx_public_phonetic_annotations_replayed_exactly": int(
            xlsx_replay["phonetic_annotations_observed"]
        )
        == int(xlsx_replay["phonetic_annotations_replayed"]),
    }
    return {
        "schema_version": "real_document_text_validation_v2",
        "generated_at": utc_now(),
        "status": "ok" if all(checks.values()) else "failed",
        "policy": {
            "official_real_documents_only": True,
            "synthetic_or_mock_records_counted": False,
            "official_public_domain_personal_information_retained": True,
            "native_text_is_not_sanitized_or_identity_stripped": True,
            "pdf_plain_text_angle_brackets_retained": True,
            "xlsx_raw_to_cas_independent_parser": str(xlsx_replay["validator"]),
            "xlsx_formula_validation_status": str(
                xlsx_replay["formula_validation_status"]
            ),
            "xlsx_public_phonetic_annotations_independently_replayed": True,
        },
        "extraction_algorithm_version": DOCUMENT_EXTRACTION_ALGORITHM_VERSION,
        "inputs": {
            "inventory": {
                "path": display_path(inventory_path),
                "sha256": sha256_path(inventory_path),
            },
            "provenance": {
                "path": display_path(provenance_path),
                "sha256": sha256_path(provenance_path),
            },
            "manifest": {
                "path": display_path(manifest_path),
                "rows": observed_rows,
                "bytes": manifest_path.stat().st_size,
                "sha256": manifest_sha256,
            },
            "extraction_report": {
                "path": display_path(extraction_report_path),
                "bytes": extraction_report_path.stat().st_size,
                "sha256": sha256_path(extraction_report_path),
            },
        },
        "validation_scope": {
            "full_raw_file_instance_checksums": bool(args.full_raw_check),
            "workers": int(args.workers),
            "batch_size": int(args.batch_size),
            "xlsx_raw_to_cas_workers": int(xlsx_replay["workers"]),
            "xlsx_raw_to_cas_fetch_batch_rows": int(
                xlsx_replay["fetch_batch_rows"]
            ),
            "xlsx_raw_to_cas_max_in_flight": int(
                xlsx_replay["max_in_flight"]
            ),
            "xlsx_raw_to_cas_validator": str(xlsx_replay["validator"]),
            "xlsx_raw_to_cas_validator_version": str(
                xlsx_replay["validator_version"]
            ),
        },
        "totals": totals,
        "diagnostics": {
            "manifest_value_mismatches": mismatch_rows,
            "invalid_work_states": invalid_states,
            "raw_error_samples": raw_errors,
            "text_error_samples": text_errors,
            "xlsx_raw_to_cas_error_samples": xlsx_replay["failures"],
        },
        "checks": checks,
        "next_action": (
            "Run bounded page-level OCR for every distinct empty or sparse PDF "
            "page and publish it as additive evidence beside native text."
        ),
    }


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if int(args.workers) < 1 or int(args.batch_size) < 1:
        print("ERROR: workers and batch-size must be positive")
        return 2
    try:
        report = validate(args)
        write_json_atomic(Path(args.report_out), report)
        print(json.dumps(report, ensure_ascii=True, sort_keys=True))
        if args.enforce and report["status"] != "ok":
            return 1
        return 0
    except (
        CorpusValidationError,
        GenerationManifestError,
        OSError,
        sqlite3.Error,
        ValueError,
        json.JSONDecodeError,
    ) as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
